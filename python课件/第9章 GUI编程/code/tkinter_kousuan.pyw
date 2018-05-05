import random
import os
import tkinter
import tkinter.ttk
from docx import Document

columnsNumber = 4

def main(rowsNumber=20, grade=4):
    if grade < 3:
        operators = '＋－'
        biggest = 20
    elif grade <= 4:
        operators = '＋－×÷'
        biggest = 100
    elif grade == 5:
        operators = '＋－×÷('
        biggest = 100
        
    document = Document()
    table = document.add_table(rows=rowsNumber, cols=columnsNumber)
    #table.style.font.name = '宋体'
    for row in range(rowsNumber):
        for col in range(columnsNumber):
            first = random.randint(1, biggest)
            second = random.randint(1, biggest)
            operator = random.choice(operators)
            
            
            if operator != '(':
                if operator == '－':
                    if first < second:
                        first, second = second, first
                r = str(first).ljust(2, ' ') + ' ' + operator  + str(second).ljust(2, ' ') + '='
            else:
                third = random.randint(1, 100)
                while True:
                    o1 = random.choice(operators)
                    o2 = random.choice(operators)
                    if o1 != '(' and o2 != '(':
                        break
                rr = random.randint(1, 100)
                if rr > 50:
                    if o2 == '－':
                        if second < third:
                            second, third = third, second
                    r = str(first).ljust(2, ' ') + o1 + ' ( ' + str(second).ljust(2, ' ') + o2 + str(third).ljust(2, ' ') + ')='
                else:
                    if o1 == '－':
                        if first < second:
                            first, second = second, first
                    r = '(' + str(first).ljust(2, ' ') + o1 + str(second).ljust(2, ' ') + ')' + o2 + str(third).ljust(2, ' ') + '='
            #cell = table.rows[row].cells[col]
            cell = table.cell(row, col)
            cell.text = r
            
    document.save('kousuan.docx')
    os.startfile('kousuan.docx')
if __name__ == '__main__':
    app = tkinter.Tk()
    app.title('KouSuan------by Dong Fuguo')
    app['width'] = 300
    app['height'] = 150
    labelNumber = tkinter.Label(app, text='Number:', justify=tkinter.RIGHT, width=50)
    labelNumber.place(x=10, y=40, width=50,height=20)
    comboNumber = tkinter.ttk.Combobox(app, values=(100,200,300,400,500), width=50)
    comboNumber.place(x=70, y=40, width=50, height=20)
    
    labelGrade = tkinter.Label(app, text='Grade:', justify=tkinter.RIGHT, width=50)
    labelGrade.place(x=130, y=40, width=50,height=20)
    comboGrade = tkinter.ttk.Combobox(app, values=(1,2,3,4,5), width=50)
    comboGrade.place(x=200, y=40, width=50, height=20)

    def generate():
        number = int(comboNumber.get())
        grade = int(comboGrade.get())
        main(number, grade)
    buttonGenerate = tkinter.Button(app, text='GO', width=40, command=generate)
    buttonGenerate.place(x=130, y=90, width=40, height=30)

    app.mainloop()
